"""Builds and validates the assignment change list from the reviewed workbook. Offline and read-only.
Usage: python -I build-change-list.py <review-workbook.xlsx> <backup-folder> <out-folder>
Decisions applied (owner, 2026-10-22): sheet 1 = Use Old App (all rows); sheets 2, 4 = Keep Production; sheet 3 = Needs Investigation (nothing added); sheet 5 = duplicates untouched.
Writes change-list.json / change-list.xlsx to the out folder. Never writes to a database."""
import json, re, sys, collections
import openpyxl
from openpyxl.styles import Font, PatternFill

wbp, backup, out = sys.argv[1:4]
norm = lambda s: re.sub(r'\s+', ' ', str(s or '').replace(' ', ' ').strip()).upper()
data = json.load(open(backup + '/data.json', encoding='utf-8'))
ws = openpyxl.load_workbook(wbp)['1 Missing assignments']
rows = [r for r in ws.iter_rows(min_row=3, values_only=True) if r[0]]

prod = collections.defaultdict(list)
for t in ('serialized_assets', 'active_laptops_np'):
    for x in data[t]: prod[norm(x['serial'])].append((t, x))

changes, problems = [], []
for r in rows:
    serial, dev, ostatus, oassignee = r[0], r[1], r[2], r[3]
    hits = prod.get(norm(serial), [])
    issues, notes = [], []
    if len(hits) != 1: issues.append(f'serial matches {len(hits)} production rows (must be exactly 1)')
    t, x = hits[0] if hits else (None, None)
    if x:
        if x['status'] != 'available': issues.append(f"production status is {x['status']}, not available")
        if (x.get('assigned_to') or '').strip(): issues.append(f"production already has assignee {x['assigned_to']}")
    name = re.sub(r'\s+', ' ', (oassignee or '').replace(' ', ' ').strip())
    if name != (oassignee or ''): notes.append('assignee had extra spaces (cleaned)')
    if not name: issues.append('empty assignee')
    if re.search(r'[&/,;]| and |\(|\?|unidentified|loaner|previous', name, re.I): issues.append(f'assignee "{name}" may not be a single person')
    if name and name == name.lower() or re.search(r'(^| )[a-z]', name): notes.append('name has a lowercase word: check capitalisation')
    rec = dict(serial=serial, production_table=t, production_id=x['id'] if x else None, device=dev, production_status=x['status'] if x else None,
               old_app_assignee=oassignee, new_assignee=name, set_status='provided', issues=issues, notes=notes)
    (problems if issues else changes).append(rec)

by_person = collections.defaultdict(list)
for c in changes + problems: by_person[norm(c['new_assignee'])].append(c['serial'])
multi = {p: s for p, s in by_person.items() if len(s) > 1}
for c in changes:
    if norm(c['new_assignee']) in multi: c['notes'].append('this person gets ' + str(len(multi[norm(c['new_assignee'])])) + ' devices: ' + ', '.join(multi[norm(c['new_assignee'])]))

summary = dict(sheet1_rows=len(rows), validated_updates=len(changes), blocked=len(problems),
               by_table=dict(collections.Counter(c['production_table'] for c in changes)),
               added=0, skipped_keep_production=5 + 4, needs_investigation=7,
               duplicates_untouched=sum(1 for k, v in prod.items() if len(v) > 1),
               people_with_multiple_devices=multi)
json.dump(dict(summary=summary, changes=changes, blocked=problems), open(out + '/change-list.json', 'w', encoding='utf-8'), indent=2)

wb = openpyxl.Workbook(); sh = wb.active; sh.title = 'Validated updates'
H = ['Serial', 'Production table', 'Device', 'Current status', 'Current assignee', 'Change: status', 'Change: assignee', 'Old-app text', 'Notes']
sh.append(H)
for c in sh[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='1F3A5F')
for c in changes: sh.append([c['serial'], c['production_table'], c['device'], c['production_status'], '(blank)', 'provided', c['new_assignee'], c['old_app_assignee'], '; '.join(c['notes'])])
if problems:
    b = wb.create_sheet('Blocked'); b.append(['Serial', 'Issues'])
    for c in problems: b.append([c['serial'], '; '.join(c['issues'])])
for i, w in enumerate([18, 20, 34, 14, 16, 14, 28, 28, 70]): sh.column_dimensions[chr(65 + i)].width = w
wb.save(out + '/change-list.xlsx')
print(json.dumps(summary, indent=1))
for c in problems: print('BLOCKED', c['serial'], c['issues'])
for c in changes:
    if c['notes']: print('NOTE', c['serial'], c['notes'])
