"""Builds the review workbook from reconciliation.json (made by compare-old-app.mjs). Offline; changes nothing else.
Usage: python -I scripts/reconcile/build-review-workbook.py <reconciliation.json> <out.xlsx>
"""
import json, re, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

r = json.load(open(sys.argv[1], encoding='utf-8'))
PLACEHOLDER = re.compile(r'unidentified|loaner|\(previous\)|unknown|spare|^n/?a$|^none$|^tbd$|\?', re.I)
OPTIONS = 'Use Old App,Keep Production,Needs Investigation'
lap = r['families']['laptops']

def fmt(s): return s if s else '(blank)'
missing, conflicts, placeholders = [], [], []
for c in lap['same'] + lap['changed']:
    o, n, d = c['old'], c['new'], c.get('diffs', {})
    if PLACEHOLDER.search(o['assigned'] or ''):
        placeholders.append((c, 'laptop'))
        continue
    if not d: continue
    pure = (o['status'] == 'provided' and n['status'] == 'available' and not n['assigned'] and 'location' not in d and set(d) <= {'status', 'assigned'})
    (missing if pure else conflicts).append(c)

wb = Workbook()
HDR = PatternFill('solid', fgColor='1F3A5F')
def sheet(title, headers, rows, widths, note):
    ws = wb.create_sheet(title)
    ws.append([note]); ws['A1'].font = Font(italic=True)
    ws.append(headers)
    for c in ws[2]: c.font = Font(bold=True, color='FFFFFF'); c.fill = HDR; c.alignment = Alignment(wrap_text=True, vertical='top')
    for row in rows: ws.append(row)
    for i, w in enumerate(widths, 1): ws.column_dimensions[chr(64 + i)].width = w
    ws.freeze_panes = 'A3'
    if rows:
        col = headers.index('Review Decision') + 1
        L = chr(64 + col)
        dv = DataValidation(type='list', formula1='"%s"' % OPTIONS, allow_blank=True)
        ws.add_data_validation(dv); dv.add(f'{L}3:{L}{len(rows) + 2}')
    return ws

RD = ['Review Decision', 'Notes']
sheet('1 Missing assignments',
      ['Serial', 'Device (old app)', 'Old-app status', 'Old-app assignee', 'Production status', 'Production assignee', 'Proposed change'] + RD,
      [[c['old']['serial'], c['old']['name'], c['old']['status'], c['old']['assigned'], c['new']['status'], fmt(c['new']['assigned']),
        f"Set status Provided, assignee {c['old']['assigned']}", '', ''] for c in missing],
      [18, 30, 12, 26, 14, 20, 46, 22, 30],
      f'{len(missing)} laptops: the old app shows them provided to a named person; production has them Available with no assignee. Nothing is applied.')

def diffs_text(d): return '; '.join(f'{k}: old "{a}" / production "{b}"' for k, (a, b) in d.items())
sheet('2 Conflicts',
      ['Serial', 'Device (old app)', 'Old-app status', 'Old-app assignee', 'Production status', 'Production assignee', 'What differs'] + RD,
      [[c['old']['serial'], c['old']['name'], c['old']['status'], fmt(c['old']['assigned']), c['new']['status'], fmt(c['new']['assigned']), diffs_text(c['diffs']), '', ''] for c in conflicts],
      [18, 30, 12, 24, 14, 24, 60, 22, 30],
      f'{len(conflicts)} laptops where old app and production disagree in a way that is not just a missing assignment. Production is not changed.')

noser = [(x, 'laptop') for x in lap['noSerialOld']] + [(x, 'monitor') for x in r['families']['monitors']['noSerialOld']] + [(x, 'other') for x in r['families']['other']['noSerialOld']]
sheet('3 No serial',
      ['Type', 'Device (old app)', 'Category', 'Old-app status', 'Old-app assignee', 'Location', 'Serial to add (fill in)'] + RD,
      [[t, x['name'], x['category'], x['status'], fmt(x['assigned']), x['location'], '', '', ''] for x, t in noser],
      [10, 34, 16, 14, 24, 18, 24, 22, 30],
      'Old-app items with no serial number: they cannot be matched to production. "Use Old App" = add to production (a serial is needed for laptops).')

sheet('4 Placeholder assignees',
      ['Serial', 'Device (old app)', 'Old-app status', 'Old-app assignee (placeholder)', 'Production status', 'Production assignee', 'Proposed change'] + RD,
      [[c['old']['serial'], c['old']['name'], c['old']['status'], c['old']['assigned'], c['new']['status'], fmt(c['new']['assigned']), 'Not a real person: a human decides what this device is (real assignee / retained / available)', '', ''] for c, _ in placeholders],
      [18, 30, 12, 28, 14, 20, 60, 22, 30],
      f'{len(placeholders)} laptops whose old-app assignee is not a person. Excluded from sheet 1 so each device appears once.')

dup = wb.create_sheet('5 Duplicate serials (info)')
dup.append(['Not for decision here. These stay as they are (never merged); resolve them in the app under Access & Roles after the release.'])
dup.append(['Serial', 'Old-app copies', 'Production copies'])
for d in lap['duplicate']: dup.append([d['key'], len(d['old']), len(d['new'])])
dup.column_dimensions['A'].width = 22

rd = wb.create_sheet('Read me', 0)
for line in ['Equipment Ledger: review workbook (read-only analysis; nothing has been applied to production)',
             'Fill "Review Decision" (drop-down: ' + OPTIONS.replace(',', ' / ') + ') and Notes on sheets 1-4.',
             'Use Old App = apply the old-app value; Keep Production = leave as is; Needs Investigation = decide later.',
             'Duplicates are listed on sheet 5 for information only and are not merged.',
             f'Counts: sheet 1 = {len(missing)}, sheet 2 = {len(conflicts)}, sheet 3 = {len(noser)}, sheet 4 = {len(placeholders)}.',
             'Contains business data: keep outside the repository.']:
    rd.append([line])
rd.column_dimensions['A'].width = 110
del wb['Sheet']
wb.save(sys.argv[2])
print('written', sys.argv[2], dict(missing=len(missing), conflicts=len(conflicts), no_serial=len(noser), placeholders=len(placeholders)))
